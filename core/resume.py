"""简历读取、解析与结构化。

解析顺序：PDF → pdfplumber 逐页取文本；DOCX → python-docx（段落 + 表格）；MD/TXT → 直读。
解析出的纯文本再交给 LLM 输出结构化 JSON，失败时降级为仅保留纯文本（打分仍可运行）。

命令行用法：python -m core.resume <简历文件路径>
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from loguru import logger

from core.config import RESUME_DIR, ensure_dirs
from core.db import create_resume
from core.models import Resume
from core.prompts import RESUME_STRUCTURE_SYSTEM, RESUME_STRUCTURE_USER

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}


def extract_text(path: Path) -> tuple[str, str]:
    """解析简历文件为纯文本，返回 (文本, 文件类型)。"""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件类型：{suffix}（仅支持 PDF / DOCX / MD / TXT）")

    if suffix == ".pdf":
        import pdfplumber

        pages: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        text = "\n".join(pages)
        return text.strip(), "pdf"

    if suffix == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip(), "docx"

    # Markdown / 纯文本：优先 UTF-8，失败回退 GBK
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="gbk", errors="ignore")
    return text.strip(), "md" if suffix in (".md", ".markdown") else "txt"


def structure_resume_text(raw_text: str) -> dict:
    """调用 LLM 把简历纯文本整理成结构化 JSON；失败返回空 dict（降级）。"""
    if not raw_text.strip():
        return {}
    from core.llm import chat_json

    try:
        data = chat_json(
            RESUME_STRUCTURE_SYSTEM,
            RESUME_STRUCTURE_USER.format(resume_text=raw_text[:12000]),
            temperature=0.1,
        )
        logger.info("简历结构化完成：skills={} projects={}",
                    len(data.get("skills") or []), len(data.get("projects") or []))
        return data
    except Exception as e:  # noqa: BLE001 - 结构化失败只降级，不阻断上传
        logger.warning("简历结构化失败，降级为纯文本模式：{}", e)
        return {}


def _safe_filename(name: str) -> str:
    """清理上传文件名，去掉路径与非法字符。"""
    base = Path(name).name
    base = re.sub(r'[\\/:*?"<>|\r\n]', "_", base).strip() or "resume"
    return base


def save_upload_bytes(filename: str, data: bytes) -> Path:
    """把上传的字节流写入 resumes/（加时间戳避免覆盖），返回保存后的路径。"""
    ensure_dirs()
    safe = _safe_filename(filename)
    stem, suffix = Path(safe).stem, Path(safe).suffix.lower()
    target = RESUME_DIR / f"{stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{suffix}"
    target.write_bytes(data)
    logger.info("简历文件已保存：{}", target)
    return target


def ingest_resume_file(
    path: Path, original_name: Optional[str] = None, active: bool = True, structure: bool = True
) -> Resume:
    """解析指定简历文件并写入数据库（active=True 时设为当前生效版本）。"""
    raw_text, file_type = extract_text(path)
    if not raw_text:
        raise ValueError("简历未解析出任何文本（可能是扫描版 PDF，请换一份可复制文本的简历）")

    structured = structure_resume_text(raw_text) if structure else {}
    resume = create_resume(
        filename=original_name or path.name,
        file_type=file_type,
        raw_text=raw_text,
        structured_json=json.dumps(structured, ensure_ascii=False) if structured else "",
        active=active,
    )
    logger.info("简历已入库：id={} 文件={} 类型={}", resume.id, resume.filename, file_type)
    return resume


def parse_structured(resume: Resume) -> dict:
    """读取简历的结构化 JSON，解析失败返回空 dict。"""
    if not resume.structured_json:
        return {}
    try:
        return json.loads(resume.structured_json)
    except json.JSONDecodeError:
        logger.warning("简历 {} 的结构化 JSON 无法解析，按纯文本处理", resume.id)
        return {}


def build_resume_brief(resume: Resume, raw_limit: int = 3500) -> str:
    """拼装给 LLM 使用的简历摘要：结构化信息优先，附带原文截断。"""
    structured = parse_structured(resume)
    lines: list[str] = []

    if structured.get("name"):
        lines.append(f"姓名：{structured['name']}")
    if structured.get("education"):
        lines.append(f"教育：{structured['education']}")
    if structured.get("skills"):
        lines.append("技能：" + "、".join(structured["skills"]))
    projects = structured.get("projects") or []
    for p in projects:
        if not isinstance(p, dict):
            continue
        parts = [p.get("name", ""), p.get("stack", "")]
        desc = "；".join(x for x in (p.get("duty", ""), p.get("result", "")) if x)
        lines.append("项目：" + " | ".join(x for x in parts if x) + (f" —— {desc}" if desc else ""))
    for exp in structured.get("experiences") or []:
        lines.append(f"经历：{exp}")
    if structured.get("highlights"):
        lines.append("亮点：" + "；".join(structured["highlights"]))

    brief = "\n".join(lines)
    brief += "\n\n【简历原文摘要】\n" + resume.raw_text[:raw_limit]
    return brief.strip()


def _main() -> int:
    """命令行入口：解析并入库一份简历文件。"""
    import sys

    from core.log import setup_logging

    setup_logging()
    if len(sys.argv) < 2:
        print("用法：python -m core.resume <简历文件路径>")
        return 2
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"文件不存在：{path}")
        return 2
    resume = ingest_resume_file(path)
    structured = parse_structured(resume)
    print(f"已入库：id={resume.id} 文件={resume.filename} 类型={resume.file_type} 字数={len(resume.raw_text)}")
    print("结构化结果：")
    print(json.dumps(structured, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
